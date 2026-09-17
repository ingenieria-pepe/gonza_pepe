"""Informe de Ventas detallado (app/modules/informes/ventas_articulos.py).

Réplica del "Estadísticas por Artículo" de Macrosoft en su vista detallada,
con la columna nueva de precio por unidad. Acá se prueba la parte PURA
(armar_detalle), el Excel, y las guardas del SQL que protegen las reglas
descubiertas con datos reales:

  - los RECIBOS son TipDoc 'V' y TAMBIÉN escriben filas en Lineas: sin el
    filtro Afecta=1 restarían importes fantasma (5.279 filas en 5 semanas).
  - Efact_Transacciones guarda reintentos: sin TOP 1 el join duplica líneas.
  - el precio por unidad es importe/cantidad (no cantidad/importe):
    18.000 / 24 cajas = 750, validado contra Lineas.Precio de la factura
    143678 de ROWANS SA.
"""
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.modules.informes import ventas_articulos as va


def _lineas_falsas():
    return [
        {"fecha": datetime(2026, 8, 13, 0, 0), "comprobante": "E-Factura Credito",
         "nro_doc": 24068, "codart": "010101-4", "descripcion": "Banana Brasil Color 4",
         "deposito": "A", "cantidad": Decimal("24.000"),
         "precio_lista": Decimal("750.0000"), "importe_sin_iva": Decimal("14754.10"),
         "iva": Decimal("3245.90"), "importe": Decimal("18000.00"),
         "cod_cliente": 43801, "cliente": "ROWANS SA", "ruc": "214660310018",
         "vendedor": "Ramiro", "tipo_cfe": "E-Factura", "efact_serie": "A",
         "efact_nro": 324260,
         "doc_cod": "02", "nrofact": 143678, "ref_doc": "", "ref_nrofact": 0},
        # La N/C: cantidad e importes negativos (sin IVA e IVA también), el
        # precio de la línea queda POSITIVO. Su línea referencia a la factura
        # por DocRef/Referencia (mismo mecanismo que la imputación de recibos).
        {"fecha": datetime(2026, 8, 14, 0, 0), "comprobante": "E-Factura N/C Caja",
         "nro_doc": 2749, "codart": "010101-4", "descripcion": "Banana Brasil Color 4",
         "deposito": "A", "cantidad": Decimal("-24.000"),
         "precio_lista": Decimal("750.0000"), "importe_sin_iva": Decimal("-14754.10"),
         "iva": Decimal("-3245.90"), "importe": Decimal("-18000.00"),
         "cod_cliente": 43801, "cliente": "ROWANS SA", "ruc": "214660310018",
         "vendedor": "Diego Cajero", "tipo_cfe": "E-Factura Nota de Credito",
         "efact_serie": "A", "efact_nro": 49459,
         "doc_cod": "204", "nrofact": 144337, "ref_doc": "02", "ref_nrofact": 143678},
        # Un descuento con cantidad 0 y sin CFE: precio vacío, número None.
        {"fecha": datetime(2026, 8, 14, 0, 0), "comprobante": "E-Ticket Contado",
         "nro_doc": 100, "codart": "D06", "descripcion": "Dto. Kiwi",
         "deposito": "", "cantidad": Decimal("0.000"),
         "precio_lista": Decimal("100.0000"), "importe_sin_iva": Decimal("-409.84"),
         "iva": Decimal("-90.16"), "importe": Decimal("-500.00"),
         "cod_cliente": 555, "cliente": "ALMACÉN", "ruc": "",
         "vendedor": "", "tipo_cfe": "E-Ticket", "efact_serie": "",
         "efact_nro": None,
         "doc_cod": "07", "nrofact": 150000, "ref_doc": "", "ref_nrofact": 0},
    ]


def test_armar_detalle_precio_por_linea_y_serializacion():
    res = va.armar_detalle(_lineas_falsas())
    f0, f1, f2 = res["filas"]

    # El precio de la línea es el precio al que se facturó, IVA incluido (750
    # en la factura y también en la N/C: -18.000 / -24). Con cantidad 0 no hay
    # precio que informar (None, no división por cero).
    assert f0["precio_unitario"] == 750.0 and f1["precio_unitario"] == 750.0
    assert f2["precio_unitario"] is None
    # sin IVA + IVA = importe, con el mismo signo en las tres.
    for f in res["filas"]:
        assert abs(f["importe_sin_iva"] + f["iva"] - f["importe"]) < 0.01
    # Fechas listas para JSON (ISO, no datetime) y columnas nuevas presentes.
    assert f0["fecha"] == "2026-08-13"
    assert f0["deposito"] == "A" and f0["precio_lista"] == 750.0
    assert f0["efact_nro"] == 324260 and f2["efact_nro"] is None

    # Los totales netean la factura contra su N/C.
    t = res["totales"]
    assert t["lineas"] == 3 and t["cantidad"] == 0.0
    assert t["importe"] == -500.0
    assert abs(t["importe_sin_iva"] + t["iva"] - t["importe"]) < 0.01


def test_armar_detalle_vacio():
    res = va.armar_detalle([])
    assert res["filas"] == []
    assert res["totales"]["lineas"] == 0 and res["totales"]["importe"] == 0


def test_guardas_del_sql():
    """Las reglas que salieron de mirar los datos reales, como asserts
    estructurales (si alguien las saca, esto lo cuenta antes que la contadora)."""
    for sql in (va.DETALLE_SQL, va.OTRAS_MONEDAS_SQL):
        # Recibos/rechazos/resguardos afuera (también son TipDoc 'V').
        assert "d.Afecta = 1" in sql
        assert "LTRIM(RTRIM(d.TipDoc)) = 'V'" in sql
        assert "ISNULL(cb.Anulada, 0) <> 1" in sql
    # Sin mezclar monedas: el informe es solo $, lo demás se avisa.
    assert "cb.Moneda = 1" in va.DETALLE_SQL
    assert "cb.Moneda <> 1" in va.OTRAS_MONEDAS_SQL
    # La venta va por Haber y la devolución por Debe.
    assert "l.CantidadHaber - l.CantidadDebe" in va.DETALLE_SQL
    # Efact_Transacciones guarda reintentos: sin el TOP 1 el join DUPLICA líneas.
    assert "TOP 1" in va.DETALLE_SQL and "OUTER APPLY" in va.DETALLE_SQL
    # Los totales del servidor aplican EXACTAMENTE las mismas guardas y los
    # mismos filtros que el detalle (si divergen, la pantalla miente).
    for guarda in ("d.Afecta = 1", "LTRIM(RTRIM(d.TipDoc)) = 'V'",
                   "ISNULL(cb.Anulada, 0) <> 1", "cb.Moneda = 1",
                   "{cliente}", "{articulo}"):
        assert guarda in va.TOTALES_SQL and guarda in va.DETALLE_SQL
    assert "{top}" in va.DETALLE_SQL


def test_filtros_cliente_y_articulo_son_puros_y_parametrizados():
    """Los filtros opcionales van como fragmentos fijos + parámetros en el
    orden del WHERE (cliente primero, artículo después). Jamás se interpola
    el valor del usuario en el SQL."""
    assert va.filtros(None, None) == ("", "", ())
    assert va.filtros(43801, None) == ("AND cb.CodCliente = %s", "", (43801,))
    assert va.filtros(None, " 010101-4 ") == ("", "AND LTRIM(RTRIM(l.CodArt)) = %s", ("010101-4",))
    assert va.filtros(43801, "010101-4") == (
        "AND cb.CodCliente = %s", "AND LTRIM(RTRIM(l.CodArt)) = %s", (43801, "010101-4"))
    assert va.filtros(None, "   ") == ("", "", ())


def test_vista_cajero_recorta_el_iva_y_no_muta_la_original():
    """El cajero no ve el desglose sin IVA / IVA (pedido de Lucas 17/08). El
    recorte es del BACK: esconder columnas en el front no es seguridad."""
    det = va.armar_detalle(_lineas_falsas())
    rec = va.vista_cajero(det)

    for f in rec["filas"]:
        assert "importe_sin_iva" not in f and "iva" not in f
        # El precio de lista tampoco (18/08): en artículos (Super) Macrosoft lo
        # carga sin IVA y desentona. El cajero ve importes finales con IVA.
        assert "precio_lista" not in f
        assert "importe" in f and "precio_unitario" in f
    assert "importe_sin_iva" not in rec["totales"] and "iva" not in rec["totales"]
    assert rec["totales"]["importe"] == det["totales"]["importe"]
    assert rec["vista"] == "cajero"
    # La original queda intacta (la contadora la recibe completa).
    assert "importe_sin_iva" in det["filas"][0] and "iva" in det["totales"]


def test_la_vista_se_decide_por_permisos():
    """caja sin permisos de administración → cajero; cualquiera de los tres
    permisos completos (o es_admin) → completa."""
    from types import SimpleNamespace as U
    from app.modules.informes.router_ventas_articulos import _vista_de

    assert _vista_de(U(es_admin=False, permisos={"caja"})) == "cajero"
    assert _vista_de(U(es_admin=False, permisos={"caja", "venta"})) == "cajero"
    assert _vista_de(U(es_admin=True, permisos=set())) == "completa"
    for p in ("admin", "administracion", "informe_ventas"):
        assert _vista_de(U(es_admin=False, permisos={p, "caja"})) == "completa"


def test_pdf_se_arma_en_formato_cajero():
    """El PDF (único export del cajero) se genera y es un PDF de verdad."""
    det = va.armar_detalle(_lineas_falsas())
    det.update(desde="2026-08-01", hasta="2026-08-14",
               otras_monedas_n=0, cliente={"cod": 43801, "nombre": "ROWANS SA"})
    contenido = va.construir_pdf(det)
    assert contenido.startswith(b"%PDF")
    assert len(contenido) > 1500


def test_excel_tiene_todas_las_columnas_y_el_total():
    det = va.armar_detalle(_lineas_falsas())
    det.update(desde="2026-08-01", hasta="2026-08-14",
               otras_monedas_n=0, cliente={"cod": 43801, "nombre": "ROWANS SA"})
    contenido = va.construir_excel(det)
    wb = load_workbook(BytesIO(contenido))
    assert wb.sheetnames == ["Detallado"]

    ws = wb["Detallado"]
    celdas = [[c.value for c in fila] for fila in ws.iter_rows()]
    plano = [v for fila in celdas for v in fila if v is not None]
    assert "ROWANS SA" in str(celdas[0][0])
    # Todas las columnas del export de Macrosoft + las nuestras.
    for col in ("Fecha", "Comprobante", "Nro Doc", "Código", "Artículo",
                "Depósito", "Cantidad", "Precio lista", "Importe sin IVA",
                "IVA", "Importe", "Precio por unidad", "Cod Cliente",
                "Cliente", "RUT", "Vendedor", "CFE", "Serie", "Nro CFE"):
        assert col in plano, f"falta la columna {col}"
    assert 324260 in plano and "214660310018" in plano
    assert "TOTAL" in plano and -500.0 in plano
    # La fecha va como fecha de Excel, no como texto; el precio de la línea
    # con cantidad 0 va vacío, no 0.
    fila_banana = next(f for f in celdas if f and f[3] == "010101-4")
    assert hasattr(fila_banana[0], "year")
    fila_dto = next(f for f in celdas if f and f[3] == "D06")
    assert fila_dto[11] is None


def test_la_referencia_de_la_nc_viaja_normalizada():
    """Regla nueva (02/09): cada línea lleva doc_cod/nrofact (su comprobante)
    y ref_doc/ref_nrofact (la factura que referencia, si es N/C). El vacío de
    Macrosoft (DocRef='' / Referencia=0) se normaliza a None, y la vista del
    cajero NO los recorta: el front los necesita para colgar la N/C debajo de
    su factura."""
    res = va.armar_detalle(_lineas_falsas())
    factura, nc, dto = res["filas"]

    assert factura["doc_cod"] == "02" and factura["nrofact"] == 143678
    assert factura["ref_doc"] is None and factura["ref_nrofact"] is None
    assert nc["ref_doc"] == "02" and nc["ref_nrofact"] == 143678
    assert dto["ref_nrofact"] is None

    rec = va.vista_cajero(res)
    assert rec["filas"][1]["ref_nrofact"] == 143678
    assert "SELECT" in va.DETALLE_SQL and "l.DocRef" in va.DETALLE_SQL


def test_ordenar_con_nc_cuelga_la_nc_detras_de_su_factura():
    """El PDF usa el MISMO armado que la pantalla (02/09): la N/C con
    referencia sale detrás de la última línea de la factura que anula; sin
    referencia (o factura fuera del listado) queda cronológica."""
    filas = va.armar_detalle(_lineas_falsas())["filas"]
    orden = va.ordenar_con_nc(filas)

    # factura (02/143678) → N/C colgada (204→02/143678) → dto suelto.
    assert [(f["doc_cod"], sub) for f, sub in orden] == [
        ("02", False), ("204", True), ("07", False)]

    # Si la factura no está en el listado, la N/C no se cuelga.
    solo_nc = [f for f in filas if f["doc_cod"] == "204"]
    assert va.ordenar_con_nc(solo_nc) == [(solo_nc[0], False)]


def test_filtro_pendientes_fragmentos_y_guardas():
    """El check "solo pendientes" (02/09): pendiente = factura A CRÉDITO
    ('S', DebeHaber 'D') cuyo total no cancelaron las imputaciones que la
    referencian (recibos y N/C, ambos documentos 'H', vía DocRef/Referencia).
    Apagado, los tres slots quedan vacíos y el SQL es el de siempre."""
    vacio = va.fragmentos_pendientes(False)
    assert vacio == {"pendientes_cte": "", "pendientes_join": "", "pendientes": ""}

    lleno = va.fragmentos_pendientes(True)
    assert "WITH pagos" in lleno["pendientes_cte"]
    # Las imputaciones son documentos 'H' (recibos Y N/C) que referencian.
    assert "LTRIM(RTRIM(dp.DebeHaber)) = 'H'" in lleno["pendientes_cte"]
    assert "lp.Referencia <> 0" in lleno["pendientes_cte"]
    # Solo deudas: crédito, documento 'D', saldo positivo con tolerancia.
    assert "LTRIM(RTRIM(d.EstadoDeCuenta)) = 'S'" in lleno["pendientes"]
    assert "LTRIM(RTRIM(d.DebeHaber)) = 'D'" in lleno["pendientes"]
    assert "> 0.05" in lleno["pendientes"]

    # Los tres slots existen en el detalle Y en los totales del servidor
    # (si divergen, la pantalla miente).
    for sql in (va.DETALLE_SQL, va.TOTALES_SQL):
        for slot in ("{pendientes_cte}", "{pendientes_join}", "{pendientes}"):
            assert slot in sql
