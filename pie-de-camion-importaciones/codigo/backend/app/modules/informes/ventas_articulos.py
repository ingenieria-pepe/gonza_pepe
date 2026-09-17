"""Informe de Ventas detallado — réplica de "Estadísticas por Artículo" de
Macrosoft (Manager SQL), vista Detallado por Comprobante.

Pedido de la contadora (14/08/2026), afinado por Lucas: UNA sola vista, el
detallado por comprobante con el máximo de columnas que Macrosoft tiene por
línea, más dos agregados que el original no trae:

  - filtro por CLIENTE con buscador, y
  - la columna "precio por unidad" = importe con impuestos / cantidad de la
    línea (verificado contra Lineas.Precio: 18.000 / 24 cajas = 750 exacto).

Lee Macrosoft PROD en vivo (solo lectura, como deudores): el espejo de Lineas
guarda una ventana de ~5 semanas y este informe es "por período" a elección.

Reglas del dato, validadas contra el export real del 14/08 (ROWANS SA, 22
líneas coincidiendo una por una):
  - venta = Documentos.TipDoc 'V' (las compras son 'C') **con Afecta = 1**:
    los RECIBOS también son TipDoc 'V' y también escriben filas en Lineas
    (5.279 en 5 semanas) — sin este filtro restarían importes fantasma. Afecta
    separa exactamente lo que entra a estadísticas (facturas, tickets,
    devoluciones, N.C., N.D.) de lo que no (recibos, rechazos 199, resguardos).
  - anuladas afuera.
  - la VENTA registra la cantidad en CantidadHaber y la devolución/N.C. en
    CantidadDebe → cantidad = Haber − Debe. Los importes llevan el signo del
    documento (DebeHaber 'H' resta), que coincide con el de la cantidad.
  - "con impuestos" = TotalLinea (= TotalSinIva + IvaLinea, verificado); las
    tres columnas van con el mismo signo, así sin_iva + iva = importe siempre.
  - moneda 1 solamente: mezclar pesos con dólares en una suma sería mentira.
    Si el rango tiene comprobantes de venta en otra moneda, se avisa aparte.
"""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook

from app.pg import fetch_one
from app.pg_mirror_sync import conn_src
from app.modules.informes.excel import FECHA, MONEDA, NEGRITA, _fecha, _tabla

# {cliente} se rellena con el filtro opcional ANTES de parametrizar (es un
# literal fijo, nunca input del usuario). El tipo de CFE es
# Documentos.Descripcion_Efact (el mismo texto del export de Macrosoft); la
# serie y el número salen de Efact_Transacciones con TOP 1 porque esa tabla
# guarda reintentos — un JOIN directo duplicaría líneas.
# ── Filtro "solo pendientes" (pedido 02/09) ────────────────────────────────
# Pendiente = factura A CRÉDITO (EstadoDeCuenta='S', DebeHaber='D') cuyo total
# no fue cancelado por las imputaciones que la referencian: las líneas de los
# RECIBOS y de las N/C llevan DocRef/Referencia contra la factura (el mismo
# mecanismo del relevamiento de cobranzas), y ambos son documentos 'H'. El
# contado nunca está pendiente (regla dura verificada: se paga 100% siempre).
# Los fragmentos se inyectan solo con el check prendido: el agregado de pagos
# es UN escaneo de Lineas (como los de deudores), no se paga si no se usa.
PENDIENTES_CTE = """
WITH pagos AS (
    SELECT LTRIM(RTRIM(lp.DocRef)) AS doc, lp.Referencia AS nrofact,
           SUM(lp.TotalLinea) AS aplicado
    FROM GestionPepe.dbo.Lineas lp
    JOIN GestionPepe.dbo.Documentos dp
         ON LTRIM(RTRIM(dp.Documento)) = LTRIM(RTRIM(lp.Documento))
    WHERE LTRIM(RTRIM(dp.TipDoc)) = 'V'
      AND LTRIM(RTRIM(dp.DebeHaber)) = 'H'
      AND lp.Referencia <> 0
      -- Poda clave: una imputación nunca es ANTERIOR a la factura que paga,
      -- así que para facturas del rango alcanza con mirar pagos desde `desde`
      -- en adelante. Sin esto el agregado escanea Lineas desde 2020 (76s).
      AND lp.Fecha >= %s
    GROUP BY LTRIM(RTRIM(lp.DocRef)), lp.Referencia
)
"""
PENDIENTES_JOIN = """
LEFT JOIN pagos pag
       ON pag.doc = LTRIM(RTRIM(cb.Documento)) AND pag.nrofact = cb.NroFact
"""
PENDIENTES_WHERE = """
  AND LTRIM(RTRIM(d.EstadoDeCuenta)) = 'S'
  AND LTRIM(RTRIM(d.DebeHaber)) = 'D'
  AND (COALESCE(cb.TotalDebe, 0) - COALESCE(cb.TotalHaber, 0)) - COALESCE(pag.aplicado, 0) > 0.05
"""


def fragmentos_pendientes(solo: bool) -> dict:
    """Los tres slots del filtro, vacíos cuando está apagado. Función pura."""
    if not solo:
        return {"pendientes_cte": "", "pendientes_join": "", "pendientes": ""}
    return {"pendientes_cte": PENDIENTES_CTE, "pendientes_join": PENDIENTES_JOIN,
            "pendientes": PENDIENTES_WHERE}


DETALLE_SQL = """
{pendientes_cte}
SELECT {top} cb.Fecha AS fecha,
       LTRIM(RTRIM(d.Detalle)) AS comprobante,
       cb.NroDoc AS nro_doc,
       LTRIM(RTRIM(l.CodArt)) AS codart,
       LTRIM(RTRIM(l.Descripcion)) AS descripcion,
       LTRIM(RTRIM(COALESCE(l.Deposito, ''))) AS deposito,
       l.CantidadHaber - l.CantidadDebe AS cantidad,
       l.Precio AS precio_lista,
       CASE WHEN LTRIM(RTRIM(d.DebeHaber)) = 'H'
            THEN -l.TotalSinIva ELSE l.TotalSinIva END AS importe_sin_iva,
       CASE WHEN LTRIM(RTRIM(d.DebeHaber)) = 'H'
            THEN -l.IvaLinea ELSE l.IvaLinea END AS iva,
       CASE WHEN LTRIM(RTRIM(d.DebeHaber)) = 'H'
            THEN -l.TotalLinea ELSE l.TotalLinea END AS importe,
       cb.CodCliente AS cod_cliente,
       LTRIM(RTRIM(COALESCE(cl.Nombre, cb.Nombre, ''))) AS cliente,
       LTRIM(RTRIM(COALESCE(cl.CIORuc, ''))) AS ruc,
       LTRIM(RTRIM(COALESCE(v.Nombre, ''))) AS vendedor,
       LTRIM(RTRIM(COALESCE(d.Descripcion_Efact, ''))) AS tipo_cfe,
       LTRIM(RTRIM(COALESCE(e.eFact_Serie, ''))) AS efact_serie,
       e.eFact_Numero AS efact_nro,
       LTRIM(RTRIM(cb.Documento)) AS doc_cod,
       cb.NroFact AS nrofact,
       LTRIM(RTRIM(COALESCE(l.DocRef, ''))) AS ref_doc,
       l.Referencia AS ref_nrofact
FROM GestionPepe.dbo.Lineas l
JOIN GestionPepe.dbo.Documentos d
     ON LTRIM(RTRIM(d.Documento)) = LTRIM(RTRIM(l.Documento))
JOIN GestionPepe.dbo.Cabezal cb
     ON LTRIM(RTRIM(cb.Documento)) = LTRIM(RTRIM(l.Documento))
    AND cb.NroFact = l.NroFact
LEFT JOIN GestionPepe.dbo.Clientes cl ON cl.CodCliente = cb.CodCliente
LEFT JOIN GestionPepe.dbo.Vendedores v ON v.Vendedor = cb.CodVendedor
OUTER APPLY (
    SELECT TOP 1 e2.eFact_Serie, e2.eFact_Numero
    FROM GestionPepe.dbo.Efact_Transacciones e2
    WHERE LTRIM(RTRIM(e2.Documento)) = LTRIM(RTRIM(cb.Documento))
      AND e2.NroFact = cb.NroFact
    ORDER BY e2.eFact_Numero DESC
) e
{pendientes_join}
WHERE LTRIM(RTRIM(d.TipDoc)) = 'V'
  AND d.Afecta = 1
  AND cb.Fecha >= %s AND cb.Fecha <= %s
  AND ISNULL(cb.Anulada, 0) <> 1
  AND cb.Moneda = 1
  {cliente}
  {articulo}
  {pendientes}
ORDER BY cb.Fecha, cb.NroFact, LTRIM(RTRIM(l.CodArt))
"""

# Los totales del rango calculados EN el servidor: la pantalla muestra las
# primeras N líneas pero los totales son de TODO el rango (pedido de la
# contadora 26/08: sin cliente, un mes son 15.000+ líneas — traerlas todas
# para mostrar 100 hacía que el proxy cortara por timeout). El Excel sigue
# bajando completo y suma en Python sobre las mismas filas que exporta.
TOTALES_SQL = """
{pendientes_cte}
SELECT COUNT(*) AS lineas,
       SUM(l.CantidadHaber - l.CantidadDebe) AS cantidad,
       SUM(CASE WHEN LTRIM(RTRIM(d.DebeHaber)) = 'H'
                THEN -l.TotalSinIva ELSE l.TotalSinIva END) AS importe_sin_iva,
       SUM(CASE WHEN LTRIM(RTRIM(d.DebeHaber)) = 'H'
                THEN -l.IvaLinea ELSE l.IvaLinea END) AS iva,
       SUM(CASE WHEN LTRIM(RTRIM(d.DebeHaber)) = 'H'
                THEN -l.TotalLinea ELSE l.TotalLinea END) AS importe
FROM GestionPepe.dbo.Lineas l
JOIN GestionPepe.dbo.Documentos d
     ON LTRIM(RTRIM(d.Documento)) = LTRIM(RTRIM(l.Documento))
JOIN GestionPepe.dbo.Cabezal cb
     ON LTRIM(RTRIM(cb.Documento)) = LTRIM(RTRIM(l.Documento))
    AND cb.NroFact = l.NroFact
{pendientes_join}
WHERE LTRIM(RTRIM(d.TipDoc)) = 'V'
  AND d.Afecta = 1
  AND cb.Fecha >= %s AND cb.Fecha <= %s
  AND ISNULL(cb.Anulada, 0) <> 1
  AND cb.Moneda = 1
  {cliente}
  {articulo}
  {pendientes}
"""

# Comprobantes de venta del rango en OTRA moneda (para avisar, no para sumar).
OTRAS_MONEDAS_SQL = """
SELECT COUNT(*) AS n
FROM GestionPepe.dbo.Cabezal cb
JOIN GestionPepe.dbo.Documentos d
     ON LTRIM(RTRIM(d.Documento)) = LTRIM(RTRIM(cb.Documento))
WHERE LTRIM(RTRIM(d.TipDoc)) = 'V'
  AND d.Afecta = 1
  AND cb.Fecha >= %s AND cb.Fecha <= %s
  AND ISNULL(cb.Anulada, 0) <> 1
  AND cb.Moneda <> 1
  {cliente}
"""


def armar_detalle(rows: list[dict]) -> dict:
    """Las líneas del export, serializadas, con el precio por unidad de CADA
    línea (importe/cantidad de esa línea = el precio al que se facturó, IVA
    incluido). Función PURA (los tests la cubren sin base)."""
    filas = []
    for r in rows:
        cantidad = float(r["cantidad"] or 0)
        importe = float(r["importe"] or 0)
        fecha = r["fecha"]
        filas.append({
            "fecha": fecha.date().isoformat() if hasattr(fecha, "date") else str(fecha or ""),
            "comprobante": (r["comprobante"] or "").strip(),
            "nro_doc": r["nro_doc"],
            "codart": r["codart"],
            "descripcion": (r["descripcion"] or "").strip(),
            "deposito": (r["deposito"] or "").strip(),
            "cantidad": round(cantidad, 3),
            "precio_lista": round(float(r["precio_lista"] or 0), 4),
            "importe_sin_iva": round(float(r["importe_sin_iva"] or 0), 2),
            "iva": round(float(r["iva"] or 0), 2),
            "importe": round(importe, 2),
            "precio_unitario": round(importe / cantidad, 2) if abs(cantidad) > 1e-9 else None,
            "cod_cliente": r["cod_cliente"],
            "cliente": (r["cliente"] or "").strip(),
            "ruc": (r["ruc"] or "").strip(),
            "vendedor": (r["vendedor"] or "").strip(),
            "tipo_cfe": (r["tipo_cfe"] or "").strip(),
            "efact_serie": (r["efact_serie"] or "").strip(),
            "efact_nro": r["efact_nro"],
            # Identidad interna del comprobante + la referencia de la línea a
            # su factura (Lineas.DocRef/Referencia — el mismo mecanismo con el
            # que los recibos imputan, verificado contra prod 02/09): con esto
            # el front cuelga cada N/C debajo de la factura que anula.
            # Referencia=0 / DocRef='' = línea sin referencia (ej. descuentos).
            "doc_cod": (r["doc_cod"] or "").strip(),
            "nrofact": r["nrofact"],
            "ref_doc": (r["ref_doc"] or "").strip() or None,
            "ref_nrofact": r["ref_nrofact"] or None,
        })
    return {
        "filas": filas,
        "totales": {
            "lineas": len(filas),
            "cantidad": round(sum(f["cantidad"] for f in filas), 3),
            "importe_sin_iva": round(sum(f["importe_sin_iva"] for f in filas), 2),
            "iva": round(sum(f["iva"] for f in filas), 2),
            "importe": round(sum(f["importe"] for f in filas), 2),
        },
    }


# Lo que el CAJERO no ve (pedido de Lucas 17-18/08): el desglose sin IVA/IVA
# es dato contable, y el precio de lista quedó afuera también — en los
# artículos (Super) Macrosoft lo carga SIN IVA y desentona con el resto (el
# cajero se queda con el precio por unidad, que es siempre con IVA).
CAMPOS_SOLO_CONTADORA = ("importe_sin_iva", "iva", "precio_lista")


def vista_cajero(datos: dict) -> dict:
    """La vista recortada para el rol de caja. Función PURA (no muta la
    original): saca el desglose de IVA y el precio de lista de las filas, y el
    desglose de los totales. El recorte va acá, en el BACK — esconder columnas
    en el front no es seguridad."""
    return {
        **datos,
        "filas": [
            {k: v for k, v in f.items() if k not in CAMPOS_SOLO_CONTADORA}
            for f in datos["filas"]
        ],
        "totales": {
            k: v for k, v in datos["totales"].items()
            if k not in CAMPOS_SOLO_CONTADORA
        },
        "vista": "cajero",
    }


def filtros(cod_cliente: int | None, articulo: str | None) -> tuple[str, str, tuple]:
    """Los dos filtros opcionales como fragmentos FIJOS de SQL (nunca input
    del usuario: el valor va parametrizado) más sus parámetros, en el orden
    en que aparecen en el WHERE. Función pura."""
    f_cli = "AND cb.CodCliente = %s" if cod_cliente else ""
    art = (articulo or "").strip()
    f_art = "AND LTRIM(RTRIM(l.CodArt)) = %s" if art else ""
    params: tuple = ((cod_cliente,) if cod_cliente else ()) + ((art,) if art else ())
    return f_cli, f_art, params


def detalle(desde: date, hasta: date, cod_cliente: int | None = None,
            articulo: str | None = None, limite: int | None = None,
            solo_pendientes: bool = False) -> dict:
    """El informe completo, leyendo Macrosoft prod (solo lectura).

    Con `limite` (la pantalla) trae las primeras N líneas y los totales de
    TODO el rango calculados en el servidor; sin límite (los exports) trae
    todas las líneas y suma en Python — así el Excel cuadra consigo mismo."""
    f_cli, f_art, extra = filtros(cod_cliente, articulo)
    pend = fragmentos_pendientes(solo_pendientes)
    # El %s del CTE de pendientes va PRIMERO (el CTE precede al SELECT).
    params: tuple = ((desde,) if solo_pendientes else ()) + (desde, hasta) + extra
    top = f"TOP {int(limite)}" if limite else ""
    with conn_src(timeout=300) as cs:
        cur = cs.cursor(as_dict=True)
        cur.execute(DETALLE_SQL.format(top=top, cliente=f_cli, articulo=f_art, **pend), params)
        rows = cur.fetchall()
        totales_sql = None
        if limite:
            cur.execute(TOTALES_SQL.format(cliente=f_cli, articulo=f_art, **pend), params)
            totales_sql = cur.fetchone() or {}
        # El aviso de otras monedas es por comprobante: no lo afecta el artículo.
        cur.execute(OTRAS_MONEDAS_SQL.format(cliente=f_cli), (desde, hasta) + extra[:1 if cod_cliente else 0])
        otras = int((cur.fetchone() or {}).get("n") or 0)

    datos = armar_detalle(rows)
    if totales_sql is not None:
        datos["totales"] = {
            "lineas": int(totales_sql.get("lineas") or 0),
            "cantidad": round(float(totales_sql.get("cantidad") or 0), 3),
            "importe_sin_iva": round(float(totales_sql.get("importe_sin_iva") or 0), 2),
            "iva": round(float(totales_sql.get("iva") or 0), 2),
            "importe": round(float(totales_sql.get("importe") or 0), 2),
        }
    datos["desde"] = desde.isoformat()
    datos["hasta"] = hasta.isoformat()
    datos["otras_monedas_n"] = otras
    datos["solo_pendientes"] = solo_pendientes
    datos["cliente"] = None
    if cod_cliente:
        # El nombre sale del espejo: es solo etiqueta, no filtra nada.
        c = fetch_one(
            "SELECT BTRIM(nombre) AS nombre FROM legacy.clientes WHERE codcliente = %s",
            (cod_cliente,),
        )
        datos["cliente"] = {"cod": cod_cliente, "nombre": (c or {}).get("nombre") or ""}
    datos["articulo"] = None
    if (articulo or "").strip():
        a = fetch_one(
            "SELECT BTRIM(descripcion) AS descripcion FROM legacy.articulos"
            " WHERE BTRIM(codarticulo) = %s", (articulo.strip(),))
        datos["articulo"] = {"cod": articulo.strip(),
                             "descripcion": (a or {}).get("descripcion") or ""}
    return datos


def ordenar_con_nc(filas: list[dict]) -> list[tuple[dict, bool]]:
    """El MISMO armado que la pantalla del cajero: cada N/C con referencia
    (ref_doc/ref_nrofact) se cuelga detrás de la última línea de la factura
    que anula; si la factura no está en el listado, la N/C queda cronológica.
    Devuelve (fila, es_nc_colgada). Función PURA."""
    ultima: dict[tuple, int] = {}
    for i, f in enumerate(filas):
        ultima[(f.get("doc_cod"), f.get("nrofact"))] = i
    colgadas: dict[int, list[dict]] = {}
    colgada_ids: set[int] = set()
    for f in filas:
        if f.get("ref_doc") and f.get("ref_nrofact"):
            idx = ultima.get((f["ref_doc"], f["ref_nrofact"]))
            if idx is not None:
                colgadas.setdefault(idx, []).append(f)
                colgada_ids.add(id(f))
    out: list[tuple[dict, bool]] = []
    for i, f in enumerate(filas):
        if id(f) not in colgada_ids:
            out.append((f, False))
        for nc in colgadas.get(i, ()):
            out.append((nc, True))
    return out


def construir_pdf(det: dict) -> bytes:
    """El PDF del informe — SIEMPRE en formato cajero (sin el desglose de
    IVA): es el export de reparto, no el contable. Los cajeros solo pueden
    bajar esto (el Excel se puede editar; el PDF no) — la contadora tiene su
    Excel completo por otro endpoint.

    A4 VERTICAL a pedido de caja (17/08): lo imprimen, y una impresora en
    vertical corta o achica un apaisado. Cada línea son DOS renglones: el
    principal con lo operativo (fecha, comprobante, artículo, cantidad,
    precio por unidad y el IMPORTE último y resaltado) y debajo uno
    secundario en gris con el resto (CFE, cliente, RUT, vendedor, depósito).
    Nada se corta; lo único que NO va es el precio de lista (decisión 18/08:
    el cajero no lo ve ni en pantalla ni acá).

    Desde el 02/09: header con el logo (calcado del PDF de pie de camión) y
    las N/C colgadas debajo de su factura con banda rosada, igual que la
    pantalla (ordenar_con_nc)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    d = vista_cajero(det)

    def n(v, dec=2):
        if v is None:
            return ""
        s = f"{v:,.{dec}f}"
        return s.replace(",", "@").replace(".", ",").replace("@", ".")

    chico = ParagraphStyle("chico", fontName="Helvetica", fontSize=7, leading=8)
    # Para la celda de fecha de las N/C colgadas: "N/C dd/mm/yy" entra en los
    # 1,5 cm de la columna solo a 6pt.
    minichico = ParagraphStyle("minichico", fontName="Helvetica", fontSize=6, leading=8)
    detallecito = ParagraphStyle("detallecito", fontName="Helvetica", fontSize=6, leading=7,
                                 textColor=colors.HexColor("#64748B"))
    titulo = ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=12, leading=15)
    meta = ParagraphStyle("meta", fontName="Helvetica", fontSize=8, leading=10,
                          textColor=colors.HexColor("#475569"))

    def ddmm(iso: str) -> str:
        return f"{iso[8:10]}/{iso[5:7]}/{iso[0:4]}"

    rango = f"{ddmm(d['desde'])} al {ddmm(d['hasta'])}"
    cliente = d.get("cliente")
    sufijo = f" · {cliente['nombre'] or cliente['cod']}" if cliente else ""
    if d.get("articulo"):
        sufijo += f" · {d['articulo']['descripcion'] or d['articulo']['cod']}"

    # "Importe" va ÚLTIMA y resaltada (negrita + banda de fondo): es EL número
    # de cada línea. Sin "P. lista" (pedido 18/08: en artículos Super viene sin
    # IVA y confunde — el precio por unidad es siempre con IVA).
    encabezado = ["Fecha", "Comprobante", "Código", "Artículo", "Cant.",
                  "P. unidad", "Importe"]
    filas_pdf = [encabezado]
    estilos = [
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("LEADING", (0, 0), (-1, -1), 8),
        ("ALIGN", (4, 0), (6, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
        ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#CBD5E1")),
        # La banda del Importe: fondo suave a lo alto de toda la tabla y la
        # cifra en negrita (imprime bien también en blanco y negro).
        ("BACKGROUND", (6, 1), (6, -1), colors.HexColor("#FEF3C7")),
        ("FONTNAME", (6, 1), (6, -1), "Helvetica-Bold"),
        ("TOPPADDING", (0, 0), (-1, -1), 1.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
    # Mismo armado que la pantalla: las N/C colgadas de su factura (02/09).
    for f, es_nc in ordenar_con_nc(d["filas"]):
        comp = f"{f['comprobante']} {f['nro_doc'] or ''}".strip()
        fecha_cel = f["fecha"][8:10] + "/" + f["fecha"][5:7] + "/" + f["fecha"][2:4]
        fila_ppal = len(filas_pdf)
        filas_pdf.append([
            Paragraph(f"<b>N/C</b>&nbsp;{fecha_cel}", minichico) if es_nc else fecha_cel,
            Paragraph(comp, chico),
            f["codart"],
            Paragraph(f["descripcion"], chico),
            n(f["cantidad"], 0),
            n(f["precio_unitario"]),
            n(f["importe"]),
        ])
        if es_nc:
            # Banda rosada en la N/C colgada (sin pisar la columna del Importe,
            # que mantiene su banda ámbar).
            estilos.append(("BACKGROUND", (0, fila_ppal), (5, fila_ppal),
                            colors.HexColor("#FFF1F2")))
        # El renglón secundario: TODO lo que no entra arriba, en gris chico,
        # a lo ancho de la hoja. Nada de la pantalla queda afuera del papel.
        partes = []
        if f["efact_nro"] is not None:
            partes.append(f"{f['tipo_cfe']} {f['efact_serie']} {f['efact_nro']}".strip())
        elif f["tipo_cfe"]:
            partes.append(f["tipo_cfe"])
        partes.append(f"{f['cliente']} ({f['cod_cliente'] or ''})".replace(" ()", ""))
        if f["ruc"]:
            partes.append(f"RUT {f['ruc']}")
        if f["vendedor"]:
            partes.append(f"Vend. {f['vendedor']}")
        if f["deposito"]:
            partes.append(f"Dep. {f['deposito']}")
        sub = len(filas_pdf)
        filas_pdf.append([Paragraph("  ·  ".join(partes), detallecito),
                          "", "", "", "", "", ""])
        # El SPAN corta hasta la penúltima columna: la banda del Importe sigue
        # de corrido también en el renglón secundario.
        estilos.append(("SPAN", (0, sub), (-2, sub)))
        estilos.append(("LINEBELOW", (0, sub), (-1, sub), 0.25, colors.HexColor("#E2E8F0")))
        if es_nc:
            estilos.append(("BACKGROUND", (0, sub), (5, sub), colors.HexColor("#FFF1F2")))

    t = d["totales"]
    filas_pdf.append(["", "", "", "TOTAL", n(t["cantidad"], 0), "",
                      n(t["importe"])])
    estilos.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
    estilos.append(("LINEABOVE", (0, -1), (-1, -1), 0.5, colors.HexColor("#94A3B8")))

    # 19 cm útiles de un A4 vertical con márgenes de 1 cm.
    anchos = [1.5, 4.6, 1.5, 5.4, 1.2, 2.2, 2.6]
    tabla = Table(filas_pdf, colWidths=[a * cm for a in anchos], repeatRows=1)
    tabla.setStyle(TableStyle(estilos))

    def pie(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 6)
        canvas.setFillColor(colors.HexColor("#94A3B8"))
        canvas.drawRightString(doc.pagesize[0] - 1 * cm, 0.6 * cm,
                               f"página {canvas.getPageNumber()}")
        canvas.restoreState()

    # ── Header con el logo, calcado del PDF de pie de camión ──
    import os
    from reportlab.lib.units import mm
    from reportlab.platypus import Image

    titulo_txt = Paragraph(
        "<para align='left'><b><font size=14 color='#1E418E'>INFORME DE VENTAS</font></b><br/>"
        f"<font size=8 color='#64748B'>Sistema Aloha · Pepe Banana · Almar S.A. · "
        f"{rango}{sufijo} · importes con impuestos · {t['lineas']} líneas · "
        f"total {n(t['importe'])}</font></para>",
        titulo,
    )
    logo_path = os.path.join(os.path.dirname(__file__), "..", "..", "assets", "logo_pepe.png")
    if os.path.exists(logo_path):
        header = Table(
            [[Image(logo_path, width=16 * mm, height=16 * mm, kind="proportional"), titulo_txt]],
            colWidths=[20 * mm, None],
        )
        header.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("LINEBELOW", (0, 0), (-1, -1), 1, colors.HexColor("#3856C1")),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
        ]))
    else:
        header = titulo_txt

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out, pagesize=A4,
        leftMargin=1 * cm, rightMargin=1 * cm, topMargin=1 * cm, bottomMargin=1.2 * cm,
        title="Informe de Ventas", author="Sistema Aloha",
    )
    doc.build([
        header,
        Spacer(1, 8),
        tabla,
    ], onFirstPage=pie, onLaterPages=pie)
    return out.getvalue()


def construir_excel(det: dict) -> bytes:
    """Una hoja con TODAS las columnas del detallado y el total al pie."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Detallado"

    rango = f"{det['desde']} al {det['hasta']}"
    cliente = det.get("cliente")
    sufijo = f" — {cliente['nombre'] or cliente['cod']}" if cliente else ""
    if det.get("articulo"):
        sufijo += f" — {det['articulo']['descripcion'] or det['articulo']['cod']}"

    filas = [
        [_fecha(f["fecha"]), f["comprobante"], f["nro_doc"], f["codart"],
         f["descripcion"], f["deposito"], f["cantidad"], f["precio_lista"],
         f["importe_sin_iva"], f["iva"], f["importe"], f["precio_unitario"],
         f["cod_cliente"], f["cliente"], f["ruc"], f["vendedor"],
         f["tipo_cfe"], f["efact_serie"], f["efact_nro"]]
        for f in det["filas"]
    ]
    r = _tabla(
        ws, f"Informe de Ventas detallado (importes con impuestos) — {rango}{sufijo}",
        ["Fecha", "Comprobante", "Nro Doc", "Código", "Artículo", "Depósito",
         "Cantidad", "Precio lista", "Importe sin IVA", "IVA", "Importe",
         "Precio por unidad", "Cod Cliente", "Cliente", "RUT", "Vendedor",
         "CFE", "Serie", "Nro CFE"],
        filas,
        formatos={1: FECHA, 7: '#,##0.000', 8: MONEDA, 9: MONEDA, 10: MONEDA,
                  11: MONEDA, 12: MONEDA},
    )
    t = det["totales"]
    ws.cell(r, 1, "TOTAL").font = NEGRITA
    for col, valor, formato in ((7, t["cantidad"], '#,##0.000'),
                                (9, t["importe_sin_iva"], MONEDA),
                                (10, t["iva"], MONEDA),
                                (11, t["importe"], MONEDA)):
        ws.cell(r, col, valor).number_format = formato
        ws.cell(r, col).font = NEGRITA

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
