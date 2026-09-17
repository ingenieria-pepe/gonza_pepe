"""Excel del Informe Diario.

Replica el workbook que hoy arma la administración a mano: una hoja por
sub-informe (MOV DE CAJA, VENTAS, RECIBOS, CHEQUES), las dinámicas, y la hoja
Informe con el resumen. Sale .xlsx (el original es .xls, formato de 2003 que
Macrosoft escribe línea por línea y por eso tarda tanto).
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TITULO = Font(bold=True, size=12)
HEADER = Font(bold=True, color="FFFFFF")
HEADER_BG = PatternFill("solid", fgColor="1F4E79")
NEGRITA = Font(bold=True)
MONEDA = '#,##0.00'
FECHA = 'dd/mm/yyyy'
BORDE_SUP = Border(top=Side(style="thin"))


def _ancho(ws, filas: list[list], minimo: int = 8, maximo: int = 42) -> None:
    if not filas:
        return
    for i in range(len(filas[0])):
        largo = max((len(str(f[i])) for f in filas if i < len(f) and f[i] is not None), default=0)
        ws.column_dimensions[get_column_letter(i + 1)].width = max(minimo, min(maximo, largo + 2))


def _tabla(ws, titulo: str, headers: list[str], filas: list[list],
           formatos: dict[int, str] | None = None, fila_inicial: int = 1) -> int:
    """Escribe título + encabezado + filas. Devuelve la fila siguiente libre."""
    r = fila_inicial
    ws.cell(r, 1, titulo).font = TITULO
    r += 2
    for c, h in enumerate(headers, start=1):
        cel = ws.cell(r, c, h)
        cel.font = HEADER
        cel.fill = HEADER_BG
        cel.alignment = Alignment(horizontal="center")
    ws.freeze_panes = ws.cell(r + 1, 1)
    r += 1
    for f in filas:
        for c, v in enumerate(f, start=1):
            cel = ws.cell(r, c, v)
            if formatos and c in formatos:
                cel.number_format = formatos[c]
        r += 1
    _ancho(ws, [headers] + filas)
    return r


def _fecha(v):
    """ISO -> datetime para que Excel lo trate como fecha y no como texto."""
    if not v:
        return None
    try:
        return datetime.fromisoformat(str(v).replace("Z", ""))
    except ValueError:
        return v


def _hoja_mov_caja(wb, mc: dict) -> None:
    ws = wb.create_sheet("MOV DE CAJA")
    _tabla(
        ws, "Movimientos de Caja",
        ["Fecha", "Cod", "Descripcion", "Numero", "Forma de Pago", "Mon",
         "Entrada", "Salida", "Concepto", "Rubro", "Descripcion", "Nº", "Caja", "Excluido"],
        # Entrada "MOD": devoluciones en negativo y movimientos internos en 0,
        # tal cual lo deja la contadora en su planilla (archivo MOD 29/07).
        [[_fecha(f["fecha"]), f["cod"], f["descripcion"], f["numero"], f["forma_pago"],
          f["moneda"], f.get("entrada_mod", f["entrada"]), f["salidas"], f["concepto"],
          f["rubro"], f["rubro_desc"], f["nro_caja"], f["caja_desc"],
          "SI" if f.get("excluido") else ""] for f in mc["filas"]],
        formatos={1: FECHA, 7: MONEDA, 8: MONEDA},
    )
    ws = wb.create_sheet("dinamica de caja")
    _tabla(ws, "Suma de Entrada", ["Forma de Pago", "Total"],
           [[d["forma"], d["total"]] for d in mc["dinamica"]] + [["Total general", mc["total"]]],
           formatos={2: MONEDA})


def _hoja_ventas(wb, v: dict) -> None:
    ws = wb.create_sheet("VENTAS")
    _tabla(
        ws, "Ventas del día",
        ["Fecha", "Documento", "Detalle", "Nro Doc", "Mon", "Importe", "Codigo",
         "Nombre", "Vendedor", "Nro Fact", "Efact Doc", "Efact Serie", "Efact Nro",
         "CONTADO/CREDITO", "Excluido"],
        [[_fecha(f["fecha"]), f["documento"], f["detalle"], f["nro_doc"], f["moneda"],
          f["importe"], f["codigo"], f["nombre"], f["vendedor"], f["nro_fact"],
          f.get("efact_doc", ""), f.get("efact_serie", ""), f.get("efact_nro"),
          f["contado_credito"], "SI" if f["excluido"] else ""] for f in v["filas"]],
        formatos={1: FECHA, 6: MONEDA},
    )


def _hoja_recibos(wb, rec: dict) -> None:
    ws = wb.create_sheet("RECIBOS")
    _tabla(
        ws, "Recibos del día",
        ["Fecha", "Doc", "Detalle", "Nro Doc", "Mon", "Importe", "Forma Pago",
         "Cod Cliente", "Nombre", "Excluido"],
        [[_fecha(f["fecha"]), f["documento"], f["detalle"], f["nro_doc"], f["moneda"],
          f["importe"], f["forma_pago"], f["cod_cliente"], f["nombre"],
          "SI" if f["excluido"] else ""] for f in rec["filas"]],
        formatos={1: FECHA, 6: MONEDA},
    )


def _hoja_cheques(wb, chq: dict) -> None:
    ws = wb.create_sheet("CHEQUES")
    _tabla(
        ws, "Documentos por Estado",
        ["Tipo", "Vence", "Recibido", "Emision", "Emitido", "Nombre Cliente",
         "Banco", "Cheque", "Mon", "Importe", "Estado", "Nro Banco", "VENCIMIENTO"],
        [[f["tipo"], _fecha(f["vence"]), _fecha(f["recibido"]), _fecha(f["emision"]),
          f["emitido"], f["nombre_cliente"], f["banco"], f["cheque"], f["moneda"],
          f["importe"], f["estado"], f["nrobanco"], f["vencimiento_grupo"]]
         for f in chq["filas"]],
        formatos={2: FECHA, 3: FECHA, 4: FECHA, 10: MONEDA},
    )
    ws = wb.create_sheet("dinamica de cheques")
    _tabla(ws, "Suma de Importe", ["VENCIMIENTO", "Total"],
           [[d["grupo"], d["total"]] for d in chq["dinamica"]]
           + [["Total general", chq["total_cartera"]]],
           formatos={2: MONEDA})


def _hoja_deudores(wb, deu: dict, fecha: str) -> None:
    ws = wb.create_sheet("DEUDORES")
    _tabla(ws, f"Deudores al {_ddmmyyyy(fecha)}  —  total {deu.get('saldo', 0):,.2f}",
           ["Cod Cliente", "Nombre", "Saldo"],
           [[d["cod_cliente"], d.get("nombre", ""), d["saldo"]] for d in deu.get("top", [])],
           formatos={3: MONEDA})


# Sub-informes que se pueden bajar por separado, para poner al lado del export
# de Macrosoft y comparar.
HOJAS_SUELTAS = {
    "mov-caja": "MOV DE CAJA",
    "ventas": "VENTAS",
    "recibos": "RECIBOS",
    "cheques": "CHEQUES",
    "deudores": "DEUDORES",
}


def construir_hoja(datos: dict, hoja: str) -> bytes:
    """Excel de UN sub-informe solo."""
    wb = Workbook()
    wb.remove(wb.active)   # arranca vacío: cada builder crea la suya
    if hoja == "mov-caja":
        _hoja_mov_caja(wb, datos["mov_caja"])
    elif hoja == "ventas":
        _hoja_ventas(wb, datos["ventas"])
    elif hoja == "recibos":
        _hoja_recibos(wb, datos["recibos"])
    elif hoja == "cheques":
        _hoja_cheques(wb, datos["cheques"])
    elif hoja == "deudores":
        # Período: el saldo es al CIERRE del rango.
        _hoja_deudores(wb, datos.get("deudores") or {},
                       datos.get("fecha_hasta") or datos["fecha"])
    else:
        raise KeyError(hoja)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def construir(datos: dict) -> bytes:
    wb = Workbook()

    # ── Hoja Informe: el resumen que hoy se completa a mano ──
    ws = wb.active
    ws.title = "Informe"
    v, rec, mc, chq, res = (datos["ventas"], datos["recibos"], datos["mov_caja"],
                            datos["cheques"], datos["resumen"])
    titulo = (
        f"INFORME  del {_ddmmyyyy(datos['fecha'])} al {_ddmmyyyy(datos['fecha_hasta'])}"
        if datos.get("fecha_hasta")
        else f"INFORME DIARIO  {_ddmmyyyy(datos['fecha'])}"
    )
    ws.cell(1, 1, titulo).font = Font(bold=True, size=14)

    bloques: list[tuple[str, list[tuple[str, float | str]]]] = [
        ("Ventas crédito / contado", [
            ("Contado", v["contado"]),
            ("Crédito", v["credito"]),
            ("Total general", v["total"]),
        ]),
        ("Ingresos del día", [(d["forma"], d["total"]) for d in mc["dinamica"]]
                            + [("Total ingresos", mc["total"])]),
        ("Cobranzas", [
            ("Recibos", rec["total"]),
            ("Cobranzas Rbos - Contado", res.get("cobranzas_mas_contado", 0)),
            ("Diferencia por apertura de caja y cierre", res.get("diferencia_caja", 0)),
        ]),
        ("Cheques", [(d["grupo"], d["total"]) for d in chq["dinamica"]]
                    + [("Cheques a acreditar", chq["a_acreditar"]),
                       ("TOTAL CHEQUES", chq["total"])]),
    ]
    deu = datos.get("deudores") or {}
    verif = deu.get("verificacion")
    if verif:
        bloques.append(("Deudores", [
            ("Saldo al inicio del período" if datos.get("fecha_hasta")
             else "Saldo del día anterior", verif["saldo_anterior"]),
            ("+ Ventas a crédito", verif["ventas_credito"]),
            ("- Cobranzas (recibos)", -verif["cobranzas"]),
            ("Saldo esperado", verif["esperado"]),
            ("Saldo deudores real", verif["real"]),
            ("Diferencia", verif["diferencia"]),
        ]))
    r = 3
    for titulo, items in bloques:
        ws.cell(r, 1, titulo).font = NEGRITA
        r += 1
        for etiqueta, valor in items:
            ws.cell(r, 1, etiqueta)
            cel = ws.cell(r, 2, valor)
            cel.number_format = MONEDA
            if str(etiqueta).lower().startswith(("total", "cheques a acreditar")):
                ws.cell(r, 1).font = NEGRITA
                cel.font = NEGRITA
                ws.cell(r, 1).border = BORDE_SUP
                cel.border = BORDE_SUP
            r += 1
        r += 1

    # Las frases que la administración escribe a mano al pie.
    ws.cell(r, 1, f"Hoy hubo una venta a Crédito de {round(v['pct_credito'] * 100)}% "
                  f"y Contado de {round(v['pct_contado'] * 100)}%")
    r += 1
    ws.cell(r, 1, f"En los ingresos del día el {round(res['pct_cobranzas'] * 100)}% corresponde "
                  f"a cobranzas y el {round(res['pct_contados'] * 100)}% de Contados")
    r += 2
    if v["excluidos_n"]:
        ws.cell(r, 1, f"Se excluyeron {v['excluidos_n']} comprobantes de supermercados "
                      f"por {v['excluido_total']:,.2f} (Devoto, Disco, Henderson, Odaler, "
                      f"Cafanor, Macromercado).").font = Font(italic=True)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18

    _hoja_mov_caja(wb, mc)
    _hoja_ventas(wb, v)
    _hoja_recibos(wb, rec)
    _hoja_cheques(wb, chq)
    if deu.get("top"):
        _hoja_deudores(wb, deu, datos.get("fecha_hasta") or datos["fecha"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ddmmyyyy(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso).strftime("%d/%m/%Y")
    except ValueError:
        return iso


def construir_saldos_mes(datos: dict) -> bytes:
    """Excel del "Reporte de Saldos por día": cliente × día del último mes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "SALDOS DIARIOS"
    fechas = datos["fechas"]
    headers = ["Cod Cliente", "Nombre"] + [_ddmm(f) for f in fechas]
    filas = [
        [f["cod_cliente"], f["nombre"]] + f["saldos"]
        for f in datos["filas"]
    ]
    _tabla(ws, f"Reporte de Saldos por día  —  total al {_ddmmyyyy(fechas[-1])}: "
               f"{datos['total_ultimo_dia']:,.2f}",
           headers, filas,
           formatos={i: MONEDA for i in range(3, len(headers) + 1)})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ddmm(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso).strftime("%d/%m")
    except ValueError:
        return iso
