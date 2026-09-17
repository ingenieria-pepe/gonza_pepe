"""Puente read-only Plan de Cargas: lee la hoja "Cargas" de los Excel maestros en
OneDrive y los espeja en ext.plan_de_cargas. Temporal, hasta el cutover a carga
nativa en Aloha.

DOS planillas, misma mecánica (cada una se activa con su URL; vacía = apagada):
  - 'BR'    → settings.plan_cargas_sync_url        (Brasil/Paraguay, la histórica)
  - 'OTROS' → settings.plan_cargas_otros_sync_url  ("PROGRAMA DE ARRIBOS": EC/CL/
              BO/PE/CO/MX y ultramar IT/GR/ES/EG — mismas columnas centrales pero
              en otro orden; ver OTROS_IDX en app.scripts.import_plan_cargas)
Cada sync REEMPLAZA solo las filas de SU fuente (no pisa a la otra).

UNA sola dirección (Excel → Aloha). NUNCA escribe en OneDrive.

Descarga: los archivos están migrados a SharePoint, así que el endpoint `shares`
da 401; lo que funciona es seguir el link con `&download=1` (cookie jar + UA de
navegador) → devuelve el .xlsx. Reusa el mapeo/columnas/upsert del importador CSV
(app.scripts.import_plan_cargas) — sólo cambia la fuente (xlsx tipado vs CSV texto).
"""
import asyncio
import datetime
import hashlib
import http.cookiejar
import io
import logging
import urllib.request
from zoneinfo import ZoneInfo

import openpyxl

from app.config import settings
from app.scripts.import_plan_cargas import fila_valida_otros, row_to_record, row_to_record_otros, upsert_rows

_UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
)
_UY = ZoneInfo("America/Montevideo")
logger = logging.getLogger(__name__)
# Debounce del poll: varias pantallas pidiendo cada 30s no disparan N descargas.
_MIN_INTERVALO_S = 25
# Cadencia del loop de fondo (server-side). La tele en kiosko NO dispara el sync
# (no autentica) → este loop la mantiene al día sin depender de ninguna pantalla.
_INTERVALO_LOOP_S = 30

FUENTES = ("BR", "OTROS")


def _estado_inicial() -> dict:
    return {
        "ultima_sync": None,   # ISO datetime UY
        "ok": None,            # True/False/None
        "filas": None,         # cargas espejadas en la última sync con cambios
        "mensaje": "todavía no sincronizó",
        "hash": None,          # sha256 del último xlsx procesado (para detectar cambios)
    }


# Estado en memoria por fuente (alcanza para un puente; al reiniciar re-sincroniza).
_ESTADOS: dict[str, dict] = {f: _estado_inicial() for f in FUENTES}

# Cómo se parsea cada planilla: fila de header (por el nombre EXACTO de la columna
# de status) y mapeo de fila → record.
_PARSERS = {
    # fila_valida: filtro extra de filas basura (None = solo el "alguna celda no
    # vacía" de siempre; la hoja OTROS tiene leyendas al pie que hay que descartar).
    "BR": {"header": "Status", "row_to_record": row_to_record, "fila_valida": None},
    "OTROS": {"header": "Status List", "row_to_record": row_to_record_otros, "fila_valida": fila_valida_otros},
}


def _url_cruda(fuente: str) -> str:
    return settings.plan_cargas_sync_url if fuente == "BR" else settings.plan_cargas_otros_sync_url


def activo(fuente: str | None = None) -> bool:
    """Sin arg: ¿hay ALGÚN puente activo? (bloquea la edición en Aloha)."""
    if fuente is not None:
        return bool(_url_cruda(fuente))
    return any(bool(_url_cruda(f)) for f in FUENTES)


def estado() -> dict:
    """Estado para el front. Top-level = fuente BR (compat con el banner viejo);
    `fuentes` = detalle por planilla (el front muestra el del tab activo)."""
    por_fuente = {
        f: {"activo": activo(f), **{k: v for k, v in _ESTADOS[f].items() if k != "hash"}}
        for f in FUENTES
    }
    return {
        "activo": activo(),
        **{k: v for k, v in _ESTADOS["BR"].items() if k != "hash"},
        "fuentes": por_fuente,
    }


def _url(fuente: str) -> str:
    u = _url_cruda(fuente)
    return u if "download=1" in u else u + ("&" if "?" in u else "?") + "download=1"


def _descargar(fuente: str) -> bytes:
    cj = http.cookiejar.CookieJar()
    op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
    req = urllib.request.Request(_url(fuente), headers={"User-Agent": _UA})
    with op.open(req, timeout=60) as r:
        data = r.read()
    if data[:2] != b"PK":  # un .xlsx es un zip → empieza con 'PK'
        raise ValueError("La descarga no es un .xlsx (¿el link dejó de ser 'cualquiera con el vínculo'?)")
    return data


def _parse(data: bytes, fuente: str) -> list[dict]:
    cfg = _PARSERS[fuente]
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    if "Cargas" not in wb.sheetnames:
        raise ValueError(f"No encontré la hoja 'Cargas' (hojas: {wb.sheetnames})")
    rows = list(wb["Cargas"].iter_rows(values_only=True))
    hdr_idx = next(
        (i for i, r in enumerate(rows)
         if r and any(isinstance(c, str) and c.strip() == cfg["header"] for c in r)),
        None,
    )
    if hdr_idx is None:
        raise ValueError(f"No encontré el header (columna '{cfg['header']}') en la hoja Cargas")
    data_rows = [
        r for r in rows[hdr_idx + 1:]
        if r and any(c is not None and str(c).strip() for c in r)
    ]
    if cfg["fila_valida"]:
        data_rows = [r for r in data_rows if cfg["fila_valida"](list(r))]
    return [cfg["row_to_record"](list(r)) for r in data_rows]


def _ahora() -> datetime.datetime:
    return datetime.datetime.now(_UY)


def _sincronizar_fuente(fuente: str, force: bool) -> None:
    est = _ESTADOS[fuente]

    # Debounce: si recién sincronizó, no vuelvas a bajar (salvo force = botón).
    if not force and est["ultima_sync"]:
        try:
            last = datetime.datetime.fromisoformat(est["ultima_sync"])
            if (_ahora() - last).total_seconds() < _MIN_INTERVALO_S:
                return
        except ValueError:
            pass

    try:
        data = _descargar(fuente)
        h = hashlib.sha256(data).hexdigest()
        if h == est["hash"] and not force:
            est.update(ultima_sync=_ahora().isoformat(), ok=True, mensaje="sin cambios")
            return
        filas = _parse(data, fuente)  # si viene vacío/roto, upsert_rows lanza y no pisa la tabla
        res = upsert_rows(filas, replace=True, fuente=fuente)
        est.update(
            ultima_sync=_ahora().isoformat(), ok=True, filas=len(filas), hash=h,
            mensaje=f"OK · {len(filas)} cargas (borradas {res['borradas']}, conservadas {res['refs']})",
        )
    except Exception as e:  # noqa: BLE001 — queremos reportar cualquier fallo en el estado
        est.update(ultima_sync=_ahora().isoformat(), ok=False, mensaje=f"Error: {e}")


def sincronizar(force: bool = False) -> dict:
    """Baja los xlsx activos y, si cambiaron (o force), reemplaza las filas de su
    fuente en ext.plan_de_cargas. Idempotente y barato cuando no hay cambios (sólo
    compara hash). Nunca tira: ante error deja el estado de esa fuente en ok=False
    y conserva los datos ya cargados."""
    for f in FUENTES:
        if activo(f):
            _sincronizar_fuente(f, force)
    return estado()


async def sync_loop() -> None:
    """Tarea de fondo (lifespan): cada 30s sincroniza los Excel maestros activos.
    Es lo que permite que la TELE en modo kiosko (con token ?k=, sin login) vea el
    Monitor de Camiones siempre al día: el servidor sincroniza solo, sin depender
    de ninguna pantalla logueada. Un solo worker uvicorn en prod, así que no se
    duplica. sincronizar() es sync (psycopg2/openpyxl) → va a un thread."""
    if not activo():
        logger.info("Puente Plan de Cargas: APAGADO (PLAN_CARGAS_SYNC_URL y _OTROS_ vacías); no arranca el loop.")
        return
    logger.info(
        "Puente Plan de Cargas: sync loop cada %ds — ACTIVO (%s).",
        _INTERVALO_LOOP_S,
        ", ".join(f for f in FUENTES if activo(f)),
    )
    while True:
        try:
            res = await asyncio.to_thread(sincronizar, False)
            for f, est in res.get("fuentes", {}).items():
                if est.get("activo") and est.get("ok") is False:
                    logger.warning("Puente Plan de Cargas [%s]: %s", f, est.get("mensaje"))
        except Exception:
            logger.exception("Puente Plan de Cargas: error inesperado en el sync loop")
        try:
            await asyncio.sleep(_INTERVALO_LOOP_S)
        except asyncio.CancelledError:
            break
